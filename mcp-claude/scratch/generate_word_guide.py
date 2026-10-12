import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_guide():
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI-Native Insurance Medallion Architecture on Microsoft Fabric")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("End-to-End Setup, Gate Architecture & Usage Guide")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(14)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(102, 102, 102)

    doc.add_paragraph() # Spacer

    # Section 1: Overview
    h1 = doc.add_heading("1. Executive Overview & Architecture Principles", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    p = doc.add_paragraph()
    p.add_run("This document provides the authoritative technical reference for the gated Medallion Architecture workflow operating on Microsoft Fabric. The solution orchestrates raw data ingestion, business logic transformation, dimensional modeling, workflow orchestration, and semantic model publishing using an interactive, single-gate conversation protocol.")

    # Core Principles List
    p_bullet1 = doc.add_paragraph(style='List Bullet')
    r = p_bullet1.add_run("Un-prefixed Notebook Naming: ")
    r.bold = True
    p_bullet1.add_run("All Medallion notebooks drop legacy prefixes (`01_`, `02_`, `03_`) and strictly follow `Bronze_<Lob>`, `Silver_<Lob>`, `Gold_<Lob>`, and `Refresh_Semantic_Model`.")

    p_bullet2 = doc.add_paragraph(style='List Bullet')
    r = p_bullet2.add_run("Source-Matching Delta Table Names: ")
    r.bold = True
    p_bullet2.add_run("Bronze and Silver Delta table names match the exact filename stems from the raw source files in OneLake (e.g. `auto_policy.csv` -> `auto_policy`).")

    p_bullet3 = doc.add_paragraph(style='List Bullet')
    r = p_bullet3.add_run("Client STTM Gold Table Names: ")
    r.bold = True
    p_bullet3.add_run("Gold star schema tables and measures are derived directly from the Data Engineer's uploaded client-style STTM (Source-to-Target Mapping) document.")

    p_bullet4 = doc.add_paragraph(style='List Bullet')
    r = p_bullet4.add_run("Fabric Orchestrator Precedes Semantic Model: ")
    r.bold = True
    p_bullet4.add_run("The Data Factory Orchestrator Pipeline (`<Lob>_Medallion_Orchestrator`) is constructed and deployed before generating the Direct Lake Semantic Model (`<Lob>_Gold_Model`).")

    p_bullet5 = doc.add_paragraph(style='List Bullet')
    r = p_bullet5.add_run("Gated Post-Deployment Execution & Error Recovery: ")
    r.bold = True
    p_bullet5.add_run("After deploying each artifact (Bronze, Silver, Gold notebooks, Orchestrator, Semantic Model), the AI explicitly asks `[y/n]` whether to execute/run now. On `y` success, execution moves seamlessly to the next stage. On `y` failure, the error log is inspected, a fix is generated upon user approval, and re-execution occurs. On `n` or `SKIP`, a courteous thank-you exit message is returned.")

    doc.add_paragraph() # Spacer

    # Section 2: Gate Architecture (G0 to G15)
    h2 = doc.add_heading("2. Complete End-to-End Gate Architecture (G0 - G15)", level=1)
    h2.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    headers = ["Gate", "Stage & Name", "User Prompt / Question", "Behavior & Response Protocol"]
    widths = [Inches(0.8), Inches(1.8), Inches(2.5), Inches(2.4)]

    for idx, text in enumerate(headers):
        hdr_cells[idx].width = widths[idx]
        p = hdr_cells[idx].paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[idx], "003366")
        set_cell_margins(hdr_cells[idx])

    gates_data = [
        ("G0", "Start Handshake", "Welcome to AI Native Data Engineering. (Checks Fabric workspace connection)", "If active, proceeds to G0b. If inactive, prompts G0a."),
        ("G0a", "Azure CLI Auth", "Fabric workspace connection is not active. Authenticate via Azure CLI? [y/n]", "y triggers `az login`. n prints thank-you message and stops execution."),
        ("G0b", "Workspace Select", "Select Workspace: press 1 for <W1>, 2 for <W2> or S to stop [1/2/.../S]", "Sets {workspace} context. S prints thank-you message and stops execution."),
        ("G1", "Source Lakehouse", "Should I list the lakehouses in {workspace}? [y/n]", "y lists available lakehouses for source selection. n prints thank-you message and stops."),
        ("G1a", "Project Action", "Select Action: 1 Start new project, 2 Resume existing, 3 Stop [1/2/3]", "1 lists raw folders. 2 lists project target lakehouses to resume. 3 prints thank-you message."),
        ("G2", "Target Folder Resolution", "Target Resolution: Enter or confirm target folder name (Suggested: <Derived_Name>): [y/n/<target_folder_name>]", "y accepts derived name; <target_folder_name> sets custom folder. Creates target folder <target_folder_name> (<target_folder_name>_LH, notebooks, orchestrator, semantic model). Target LH creation deferred. n stops."),
        ("G2a", "LOB Skill Draft", "No LOB skill found for folder. Create lob/<lob>.md from template? [y/n]", "y drafts LOB config. n prints thank-you message and stops."),
        ("G2b", "File Details", "Target folder created (Target LH created on execution). Show file details in folder <folder>? [y/n]", "y displays file sizes, partition paths, and table stem mappings. n prints thank-you message."),
        ("G4", "Data Profiling", "Run data profiling on <LOB> raw tables? [y/n]", "y performs read-only profiling (types, nulls, distincts, candidate keys). n prints thank-you message."),
        ("G5", "Profile Approval", "Are you satisfied with the data profile? [APPROVE/DECLINE]", "APPROVE saves verified quirks to lob/<lob>.md. DECLINE triggers interactive issues loop."),
        ("G6", "Bronze Notebook", "Generate Bronze notebook for <LOB> into schema {bronze_schema}? [y/n]", "y generates `Bronze_<Lob>` (table names match raw source stems). n prints thank-you message."),
        ("G7", "DE Approve Bronze", "DE review - approve Bronze notebook? [APPROVE/DECLINE]", "APPROVE deploys `Bronze_<Lob>` to Fabric. DECLINE triggers issues loop."),
        ("G7a", "Run Bronze Prompt", "Run Bronze notebook now? [y/n]", "y executes `Bronze_<Lob>` (seamless to G8 on success; failure recovery on error). n prints thank-you message."),
        ("G8", "Silver Rules Upload", "Upload business rules for Silver [UPLOAD/SKIP]", "UPLOAD parses and merges rules (proceeds to G9). SKIP prints thank-you message and stops."),
        ("G9", "Silver Notebook", "Generate Silver notebook for <LOB>? [y/n]", "y generates `Silver_<Lob>` (source stem table names). n prints thank-you message and stops."),
        ("G10", "DE Approve Silver", "DE review - approve Silver notebook? [APPROVE/DECLINE]", "APPROVE deploys `Silver_<Lob>` to Fabric. DECLINE triggers issues loop."),
        ("G10a", "Run Silver Prompt", "Run Silver notebook now? [y/n]", "y executes `Silver_<Lob>` (seamless to G11 on success; failure recovery on error). n prints thank-you message."),
        ("G11", "Gold STTM Upload", "Upload client-style STTM document for Gold modeling [UPLOAD/SKIP]", "UPLOAD parses STTM for Gold table names/rules (proceeds to G12). SKIP prints thank-you message."),
        ("G12", "Gold Model Approve", "Approve Gold model and generate Gold notebook? [APPROVE/DECLINE]", "APPROVE generates `Gold_<Lob>` from STTM specs. DECLINE triggers model issues loop."),
        ("G12b", "DE Approve Gold", "DE review - approve Gold notebook? [APPROVE/DECLINE]", "APPROVE deploys `Gold_<Lob>` to Fabric. DECLINE triggers issues loop."),
        ("G12c", "Run Gold Prompt", "Run Gold notebook now? [y/n]", "y executes `Gold_<Lob>` (seamless to G13 on success; failure recovery on error). n prints thank-you message."),
        ("G13", "Orchestrator Create", "Create orchestrator pipeline (<Lob>_Medallion_Orchestrator)? [y/n]", "y prompts G13a/b, deploys Orchestrator Pipeline to Fabric. n prints thank-you message and stops."),
        ("G13a/b", "Pipeline Pick LH", "Source Lakehouse? / Target Lakehouse? [1/2/...]", "Sets source raw lakehouse and target medallion lakehouse for Data Factory pipeline."),
        ("G13c", "Run Orchestrator", "Run orchestrator now? [y/n]", "y executes Orchestrator Pipeline in Fabric (seamless to G14 on success). n prints thank-you message."),
        ("G14", "Semantic Model", "Generate Direct Lake semantic model on <LOB> Gold? [y/n]", "y builds Direct Lake semantic model `<Lob>_Gold_Model` from Gold tables. n prints thank-you message."),
        ("G14b", "DE Approve Model", "Approve semantic model? [APPROVE/DECLINE]", "APPROVE deploys `<Lob>_Gold_Model` in Fabric workspace. DECLINE triggers issues loop."),
        ("G14c", "Refresh Model Prompt","Refresh semantic model dataset now? [y/n]", "y triggers Direct Lake dataset refresh (seamless to G15 on success). n prints thank-you message."),
        ("G15", "Close & Summary", "Process another LOB? [y/n]", "Displays summary matrix of tables, row counts, and deployed items. y resets to G2. n stops.")
    ]

    for gate_id, stage_name, prompt_text, behavior_text in gates_data:
        row_cells = table.add_row().cells
        row_cells[0].width = widths[0]
        row_cells[1].width = widths[1]
        row_cells[2].width = widths[2]
        row_cells[3].width = widths[3]

        row_cells[0].paragraphs[0].add_run(gate_id).bold = True
        row_cells[1].paragraphs[0].add_run(stage_name)
        row_cells[2].paragraphs[0].add_run(prompt_text)
        row_cells[3].paragraphs[0].add_run(behavior_text)

        for cell in row_cells:
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

    doc.add_paragraph() # Spacer

    # Section 3: Execution, Exit & Failure Protocols
    h3 = doc.add_heading("3. Execution, Error Recovery & Thank-You Exit Protocols", level=1)
    h3.runs[0].font.color.rgb = RGBColor(0, 51, 102)

    p_exec = doc.add_paragraph()
    p_exec.add_run("1. Post-Deployment Execution Flow: ").bold = True
    p_exec.add_run("Following DE approval and deployment of any Fabric artifact (`Bronze_<Lob>`, `Silver_<Lob>`, `Gold_<Lob>`, Orchestrator Pipeline, or Semantic Model), the AI presents a explicit `[y/n]` run prompt.")

    p_succ = doc.add_paragraph(style='List Bullet')
    p_succ.add_run("On User Response 'y' (Success): ").bold = True
    p_succ.add_run("The artifact is executed via Fabric API/Livy session. Upon successful completion, row counts and status are reported, and execution transitions seamlessly to the next gate.")

    p_fail = doc.add_paragraph(style='List Bullet')
    p_fail.add_run("On User Response 'y' (Failure): ").bold = True
    p_fail.add_run("The execution failure log is extracted from `Files/_logs/<stage>_<run_id>.log` and analyzed against `failure_registry.md`. The AI presents: ")
    p_fail.add_run("\"Execution failed due to <error>. Generate updated version with fix? [y/n]\"").italic = True
    p_fail.add_run(". On `y`, the updated artifact code is generated, presented for DE approval, deployed, and re-prompted for execution. On `n`, a thank-you exit message is printed and execution terminates.")

    p_exit = doc.add_paragraph(style='List Bullet')
    p_exit.add_run("On User Response 'n' or 'SKIP': ").bold = True
    p_exit.add_run("If the user declines to run an artifact (`n`) or chooses to `SKIP` an upload gate (Silver rules or Gold STTM), the AI acknowledges the decision with a courteous thank-you message (e.g. \"Thank you! Artifact deployed to Fabric. Execution skipped as requested.\") and exits cleanly without error.")

    # Save Document
    target_path = r"c:\Users\nikhi\Downloads\mcp-claude\MCP_Claude_Setup_And_Usage_Guide.docx"
    doc.save(target_path)
    print(f"Guide successfully generated and saved to {target_path}")

if __name__ == "__main__":
    generate_guide()
